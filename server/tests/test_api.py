import os

import pytest
from fastapi.testclient import TestClient


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    d = tmp_path_factory.mktemp("data")
    os.environ["SBT_DATA_DIR"] = str(d)
    import importlib

    from sbt import config
    importlib.reload(config)
    from sbt.data import build, synthetic
    synthetic.write(d / "demo", n_codes=40, start="2018-01-01", end="2020-12-31")
    build.build(d / "demo", log=lambda *a: None)
    from sbt.api import app as app_mod
    importlib.reload(app_mod)
    return TestClient(app_mod.app)


TEXT = """[신호]
거래량 >= 거래량(1) * 5 AND 등락률 >= 10
[매수]
지정가 = (시가 + 종가) / 2
유효기간 = 5
[청산]
목표 = 3
보유기간 = 10
[대상]
시장 = 코스피, 코스닥
기간 = 2018-01-01 ~ 오늘
"""


def test_run_and_chart(client):
    r = client.post("/api/run", json={"dataset": "demo", "text": TEXT})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["kind"] == "single" and j["summary"]["n_signals"] > 0
    assert any("가상" in c["text"] for c in j["checks"])
    code = j["trades"][0]["code"]
    c = client.get("/api/chart", params={"dataset": "demo", "code": code})
    assert c.status_code == 200 and len(c.json()["bars"]) > 100


def test_sweep(client):
    r = client.post("/api/run", json={"dataset": "demo", "text": TEXT.replace("목표 = 3", "목표 = [3, 5]")})
    assert r.status_code == 200 and len(r.json()["variants"]) == 2


def test_errors(client):
    r = client.post("/api/run", json={"dataset": "demo", "text": "[신호]\n등락률 >= 12%"})
    assert r.status_code == 422 and "%" in r.json()["detail"]
    r = client.post("/api/check-expr", json={"expr": "종가 > 이평(종가)"})
    assert r.json()["ok"] is False
    r = client.post("/api/run", json={"dataset": "kr", "text": TEXT})
    assert r.status_code == 409


def test_strategies(client):
    r = client.post("/api/strategies", json={"name": "급등 눌림", "text": TEXT})
    sid = r.json()["id"]
    assert any(s["id"] == sid for s in client.get("/api/strategies").json())
    client.delete(f"/api/strategies/{sid}")
    assert not any(s["id"] == sid for s in client.get("/api/strategies").json())


def test_scan(client):
    r = client.post("/api/scan", json={"dataset": "demo", "text": "[신호]\n등락률 >= 5\n"})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["days"] == 5  # 지정가 기본 유효기간만큼
    assert j["since"] <= j["as_of"]
    for x in j["rows"]:
        assert j["since"] <= x["signal_date"] <= j["as_of"]
        assert x["change"] is None or x["change"] >= 5 - 1e-6
        assert x["kind"] in ("pending", "holding", "closed", "expired")
        if x["kind"] == "pending":
            assert x["limit"] is not None and x["fill_date"] is None
    r = client.post("/api/scan", json={"dataset": "demo", "text": "[신호]\n등락률 >= 5\n[매수]\n다음날시가\n"})
    assert r.json()["days"] == 1
    r = client.post("/api/scan", json={"dataset": "demo", "text": "[신호]\n등락률 >= [5, 6]\n"})
    assert r.status_code == 422


def test_explore(client):
    import time as _t
    text = "[신호]\n등락률 >= [5, 8]\n[청산]\n목표 = [3, 5]\n"
    r = client.post("/api/explore/plan", json={"dataset": "demo", "text": text})
    assert r.status_code == 200, r.text
    assert r.json()["n_combos"] == 4 and r.json()["user_defined"]
    r = client.post("/api/explore", json={"dataset": "demo", "text": text, "split": "2020-01-01"})
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    for _ in range(120):
        j = client.get(f"/api/explore/{jid}").json()
        if j["status"] != "running":
            break
        _t.sleep(0.5)
    assert j["status"] == "done", j
    res = j["result"]
    assert res["n_total"] == 4 and res["split"] == "2020-01-01"
    for row in res["passed"] + res["near_miss"]:
        assert row["is"]["n_complete"] + row["oos"]["n_complete"] == row["all"]["n_complete"]
    # 자동 범위: 신호 기준값 ±30%, 청산 후보
    r = client.post("/api/explore/plan", json={"dataset": "demo", "text": "[신호]\n등락률 >= 10\n"})
    axes = r.json()["axes"]
    assert axes[0]["values"] == [7.0, 10.0, 13.0]
    assert r.json()["n_combos"] == 3 * 27


def test_journal(client):
    from sbt import journal as J
    costs = {"buy_fee_pct": 0.0, "sell_fee_pct": 0.0, "sell_tax_pct": 0.0}
    # 분할 매수 2번(100@10, 100@12) → 평균 11, 분할 매도 50@15, 150@9
    fills = [{"side": "buy", "date": "2020-01-02", "price": 10, "qty": 100},
             {"side": "buy", "date": "2020-01-03", "price": 12, "qty": 100},
             {"side": "sell", "date": "2020-01-06", "price": 15, "qty": 50},
             {"side": "sell", "date": "2020-01-08", "price": 9, "qty": 150}]
    p = J.position(fills, costs, None)
    assert p["avg_buy"] == 11 and p["closed"]
    assert p["realized"] == (15 - 11) * 50 + (9 - 11) * 150  # 200 - 300 = -100
    assert p["ret"] == round(-100 / 2200 * 100, 2)
    p2 = J.position(fills[:3], costs, 13.0)  # 150주 보유 중, 현재가 13
    assert p2["status"] == "open" and p2["unrealized"] == (13 - 11) * 150 and p2["realized"] == 200
    import pytest as _pt
    with _pt.raises(J.JournalError):
        J.position(fills[:1] + [{"side": "sell", "date": "2020-01-03", "price": 11, "qty": 101}], costs, None)

    # API: 신호 종목 하나로 기록 → 전략 비교까지
    scan = client.post("/api/scan", json={"dataset": "demo", "text": "[신호]\n등락률 >= 3\n", "days": 20}).json()
    row = next(r for r in scan["rows"] if r["fill_date"])
    body = {"dataset": "demo", "code": row["code"], "signal_date": row["signal_date"],
            "strategy_name": "테스트", "strategy_text": "[신호]\n등락률 >= 3\n",
            "fills": [{"side": "buy", "date": row["fill_date"], "price": row["fill_price"], "qty": 10}]}
    r = client.post("/api/journal", json=body)
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["name"] == row["name"] and t["pos"]["status"] == "open"
    assert t["compare"]["strategy"]["fill_date"] == row["fill_date"]
    r = client.post(f"/api/journal/{t['id']}/fills", json={"side": "sell", "date": row["fill_date"], "price": 1, "qty": 11})
    assert r.status_code == 422  # 보유 수량보다 많이
    r = client.post(f"/api/journal/{t['id']}/fills", json={"side": "sell", "date": t["last_date"], "price": t["last_close"], "qty": 10})
    assert r.status_code == 200 and r.json()["pos"]["closed"]
    st = client.get("/api/journal/stats", params={"dataset": "demo"}).json()
    assert st["closed"]["n"] == 1 and st["by_strategy"][0]["name"] == "테스트"
    assert client.get("/api/stocks", params={"dataset": "demo", "q": row["code"][:4]}).json()
    x = client.get("/api/journal/export", params={"dataset": "demo"})
    assert x.status_code == 200 and x.content[:2] == b"PK" and "attachment" in x.headers["content-disposition"]
    import io as _io
    import openpyxl as _ox
    wb = _ox.load_workbook(_io.BytesIO(x.content))
    assert wb.sheetnames == ["거래", "체결", "월별", "전략별"] and wb["체결"].max_row == 3
    assert client.delete(f"/api/journal/{t['id']}").json()["ok"]
    assert client.get("/api/journal", params={"dataset": "demo"}).json() == []
